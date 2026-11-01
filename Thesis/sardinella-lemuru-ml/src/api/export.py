"""Session-history filtering and evidence-aware CSV/XLSX export helpers."""

from __future__ import annotations

import csv
from datetime import date, datetime
from io import BytesIO, StringIO
import json
from typing import Iterable, Mapping


# The field order is intentionally stable: it preserves the established first
# columns while adding enough source data to reproduce the final AI verdict.
EXPORT_FIELDS = {
    "fish_id": "Fish ID", "timestamp": "Timestamp", "grade": "AI Final Grade",
    "final_weighted_score": "AI Final Support (%)", "best_evidence_class": "Best Evidence Class",
    "best_evidence_score": "Best Evidence Support (%)", "verdict_status": "Verdict Status",
    "verdict_reason_code": "Verdict Reason Code", "reason_codes": "All Verdict Reason Codes", "verdict_reason": "Verdict Reason",
    "second_grade": "Second Evidence Grade", "second_support": "Second Evidence Support (%)", "grade_margin": "Top-vs-Second Margin (%)",
    "manual_grade": "Manual Grade", "manual_override": "Manual Override", "review_timestamp": "Review Timestamp",
    "reviewer": "Reviewer", "review_source": "Review Source",
    "detection_confidence": "Model 1 Detection Confidence (%)",
    "grade_confidence": "Accepted Final Grade Support (%)", "observed_region_count": "Observed Region Count",
    "observed_regions": "Observed Regions", "original_weight_coverage": "Original Evidence Coverage (%)",
    "evidence_completeness": "Evidence Completeness", "model2_observation_count": "Model 2 Usable Frame Observations",
    "model2_candidate_grading_frames": "Model 2 Candidate Grading Frames",
    "model2_inference_interval_frames": "Model 2 Inference Interval (Processed Frames)",
    "track_lifetime_detection_frames": "Track Lifetime Model 1 Detection Frames",
    "sharpness_score": "Latest Crop Sharpness (Laplacian Variance)", "best_frame_score": "Best Frame Score",
    "best_frame_id": "Best Frame ID", "best_crop_path": "Best Crop Path",
    "body_observed": "Body Observed", "body_presence": "Body Observation Status", "body_grade": "Body Top Grade",
    "body_grade_confidence": "Body Evidence Score (%)", "body_weight": "Body Original Weight",
    "body_effective_weight": "Body Effective Weight", "body_weighted_contribution": "Body Winning Contribution (%)",
    "body_class_a_contribution": "Body Class A Contribution (%)", "body_class_b_contribution": "Body Class B Contribution (%)",
    "body_class_c_contribution": "Body Class C Contribution (%)", "body_rejected_contribution": "Body Rejected Contribution (%)",
    "body_observation_count": "Body Observation Count", "body_mean_confidence": "Body Mean Evidence (%)",
    "body_max_confidence": "Body Max Evidence (%)", "body_std_dev": "Body Evidence Std Dev (%)",
    "head_observed": "Head Observed", "head_presence": "Head Observation Status", "head_grade": "Head Top Grade",
    "head_grade_confidence": "Head Evidence Score (%)", "head_weight": "Head Original Weight",
    "head_effective_weight": "Head Effective Weight", "head_weighted_contribution": "Head Winning Contribution (%)",
    "head_class_a_contribution": "Head Class A Contribution (%)", "head_class_b_contribution": "Head Class B Contribution (%)",
    "head_class_c_contribution": "Head Class C Contribution (%)", "head_rejected_contribution": "Head Rejected Contribution (%)",
    "head_observation_count": "Head Observation Count", "head_mean_confidence": "Head Mean Evidence (%)",
    "head_max_confidence": "Head Max Evidence (%)", "head_std_dev": "Head Evidence Std Dev (%)",
    "tail_observed": "Tail Observed", "tail_presence": "Tail Observation Status", "tail_grade": "Tail Top Grade",
    "tail_grade_confidence": "Tail Evidence Score (%)", "tail_weight": "Tail Original Weight",
    "tail_effective_weight": "Tail Effective Weight", "tail_weighted_contribution": "Tail Winning Contribution (%)",
    "tail_class_a_contribution": "Tail Class A Contribution (%)", "tail_class_b_contribution": "Tail Class B Contribution (%)",
    "tail_class_c_contribution": "Tail Class C Contribution (%)", "tail_rejected_contribution": "Tail Rejected Contribution (%)",
    "tail_observation_count": "Tail Observation Count", "tail_mean_confidence": "Tail Mean Evidence (%)",
    "tail_max_confidence": "Tail Max Evidence (%)", "tail_std_dev": "Tail Evidence Std Dev (%)",
    "class_a_weighted_score": "Class A Final Competing Support (%)", "class_b_weighted_score": "Class B Final Competing Support (%)",
    "class_c_weighted_score": "Class C Final Competing Support (%)", "rejected_weighted_score": "Rejected Final Competing Support (%)",
    "body_mean_hue": "Body Mean Hue (deg)", "body_median_hue": "Body Median Hue (deg)",
    "body_mean_saturation": "Body Mean Saturation (%)", "body_mean_value": "Body Mean Value (%)", "body_yellow_ratio": "Body Yellow Proxy (%)",
    "head_mean_hue": "Head Mean Hue (deg)", "head_median_hue": "Head Median Hue (deg)",
    "head_mean_saturation": "Head Mean Saturation (%)", "head_mean_value": "Head Mean Value (%)", "head_yellow_ratio": "Head Yellow Proxy (%)",
    "tail_mean_hue": "Tail Mean Hue (deg)", "tail_median_hue": "Tail Median Hue (deg)",
    "tail_mean_saturation": "Tail Mean Saturation (%)", "tail_mean_value": "Tail Mean Value (%)", "tail_yellow_ratio": "Tail Yellow Proxy (%)",
    "whole_fish_mean_hue": "Whole Fish Mean Hue (deg)", "whole_fish_median_hue": "Whole Fish Median Hue (deg)",
    "whole_fish_mean_saturation": "Whole Fish Mean Saturation (%)", "whole_fish_mean_value": "Whole Fish Mean Value (%)", "whole_fish_yellow_ratio": "Whole Fish Yellow Proxy (%)",
    "supporting_image_measurements": "Supporting Image Measurements (not used for final grade)",
    "surface_damage_status": "Surface Damage Status (supporting evidence)",
    "surface_damage_severity": "Surface Damage Severity (supporting evidence)",
    "surface_damage_confidence": "Surface Damage Confidence (%)",
    "surface_damage_length_px": "Surface Damage Estimated Length (px)",
    "surface_damage_locations": "Surface Damage Locations (JSON)",
    "surface_damage_mask_source": "Surface Damage Mask Source",
    "detected_regions": "Raw Model 2 Detected Regions", "model2_classes": "Raw Model 2 Classes",
    "adjustment_applied": "Adjustment Applied", "adjustment_reason": "Adjustment Reason",
    "override_applied": "Rejected Override Applied", "override_reason": "Rejected Override Reason",
    "model1_version": "Model 1 Checkpoint", "model1_hash": "Model 1 SHA256 (short)",
    "model2_version": "Model 2 Checkpoint", "model2_hash": "Model 2 SHA256 (short)",
    "grading_config_version": "Grading Config Version", "detection_threshold": "Model 1 Detector Threshold",
    "model2_detector_threshold": "Model 2 Detector Threshold", "part_evidence_threshold": "Part Evidence Threshold",
    "quality_threshold": "Legacy Model 2 Detector Threshold", "final_verdict_threshold": "Final Verdict Threshold",
    "minimum_coverage_setting": "Minimum Coverage Setting", "minimum_grade_margin_setting": "Minimum Grade Margin",
    "grading_mode": "Grading Mode", "quality_inference_mode": "Model 2 Inference Mode", "roi_padding": "Model 2 ROI Padding (px)",
    "tracker_backend": "Tracker Backend", "camera_settings_session_reference": "Camera Settings Session Reference",
    "final_explanation": "Calculation / Explanation", "processing_time": "Processing Time (ms)",
    "decision_latency_ms": "First Detection to Count-line Decision (ms)",
    "decision_deadline_seconds": "Configured Decision Deadline (s)",
    "decision_deadline_status": "Decision Deadline Status",
    "decision_timing_scope": "Decision Timing Scope",
    # Retain every aggregated Model 2 region/grade support value so a CSV/XLSX
    # export can later be replayed with different weights or acceptance rules.
    # These are evidence/support scores, not calibrated class probabilities.
    "body_presence_confidence": "Body Presence Evidence (%)",
    "body_class_a_evidence": "Body Class A Evidence (%)", "body_class_b_evidence": "Body Class B Evidence (%)",
    "body_class_c_evidence": "Body Class C Evidence (%)", "body_rejected_evidence": "Body Rejected Evidence (%)",
    "head_presence_confidence": "Head Presence Evidence (%)",
    "head_class_a_evidence": "Head Class A Evidence (%)", "head_class_b_evidence": "Head Class B Evidence (%)",
    "head_class_c_evidence": "Head Class C Evidence (%)", "head_rejected_evidence": "Head Rejected Evidence (%)",
    "tail_presence_confidence": "Tail Presence Evidence (%)",
    "tail_class_a_evidence": "Tail Class A Evidence (%)", "tail_class_b_evidence": "Tail Class B Evidence (%)",
    "tail_class_c_evidence": "Tail Class C Evidence (%)", "tail_rejected_evidence": "Tail Rejected Evidence (%)",
    "model2_available": "Model 2 Available", "session_id": "Session ID",
    "weighting_method": "Weighting Method", "tail_evaluation_status": "Tail Evaluation Status",
    "whole_fish_bbox_width_px": "Whole-fish Bounding-box Width (px)",
    "whole_fish_bbox_height_px": "Whole-fish Bounding-box Height (px)",
    "body_bbox_width_px": "Body Bounding-box Width (px)",
    "body_bbox_height_px": "Body Bounding-box Height (px)",
    "relative_body_width_proxy": "Relative Body Width (bounding-box proxy)",
    "body_width_initial_proxy": "Initial Relative Body Width Proxy",
    "body_width_measurement_method": "Body Width Measurement Method",
    "body_width_orientation": "Body Width Orientation",
    "body_width_reliability": "Body Width Measurement Reliability",
    "body_width_reliability_reasons": "Body Width Reliability Reasons",
    "body_width_normal_reference_range": "Body Width Normal Reference Range",
    "body_width_abnormality": "Body Width Abnormality",
    "body_width_score_status": "Body Width Score Status",
    "body_width_class_c_points": "Body Width Class C Supporting Points (0-10)",
    "body_width_class_c_points_display": "Body Width Class C Points / Calibration Status",
    "body_width_scoring_enabled": "Body Width Scoring Enabled",
    "body_width_score_contributed": "Body Width Score Contributed to Final Grade",
    "body_width_applied_body_score_shift": "Body Width Applied Body Score Shift",
    "body_width_contribution_reason": "Body Width Contribution Reason",
    "body_width_threshold_configuration": "Body Width Threshold Configuration (JSON)",
    "body_width_analysis": "Body Width Analysis (JSON; bounding-box proxy)",
    "fish_length_px": "Measured Fish Length (px)",
    "max_body_width_px": "Maximum Body Width (px)",
    "relative_fish_body_width_ratio": "Relative Body Width Ratio (segmentation)",
    "fish_geometry_status": "Fish Geometry Measurement Status",
    "fish_geometry_reliability": "Fish Geometry Measurement Reliability",
    "fish_geometry_reasons": "Fish Geometry Reliability Reasons",
    "fish_geometry_method": "Fish Geometry Measurement Method",
    "fish_geometry_orientation_degrees": "Fish Geometry Orientation (degrees)",
    "fish_geometry_measurement_validated": "Fish Geometry Anatomical Validation",
    "fish_geometry_used_for_grading": "Fish Geometry Used for Grading",
    "fish_length_cm": "Physical Fish Length (cm; calibrated)",
    "max_body_width_cm": "Physical Maximum Body Width (cm; calibrated)",
    "physical_calibration_status": "Physical Calibration Status",
    "physical_calibration": "Physical Calibration (JSON)",
    "fish_geometry_analysis": "Fish Geometry Analysis (JSON)",
    "body_width_measurement_status": "Body Width Measurement Status",
    "body_width_calibration_status": "Body Width Calibration Status",
    "body_width_measurement_available": "Body Width Dimensions Available",
    "body_width_reliable_for_scoring": "Body Width Reliable for Scoring",
    "body_width_score_applied": "Width Score Applied to Classification Evidence",
    "body_width_scoring_method": "Width Scoring Method",
    "body_width_applied_geometric_score_shift": "Width Applied Geometric Support Shift (fraction)",
    "body_width_relative_proxy_source": "Relative Body Width Proxy Measurement Source",
    "body_width_configuration_source": "Body Width Configuration Provenance",
    "body_width_roi_proxy": "Relative Body Width from Existing Pixel ROIs (diagnostic proxy)",
    "body_width_detector_proxy": "Relative Body Width from Detector Boxes (diagnostic proxy)",
    "whole_fish_roi_width_px": "Whole-fish Existing ROI Width (px)",
    "whole_fish_roi_height_px": "Whole-fish Existing ROI Height (px)",
    "body_roi_width_px": "Body Existing ROI Width (px)",
    "body_roi_height_px": "Body Existing ROI Height (px)",
    "original_weighted_grade": "Original Weighted Grade",
    "original_weighted_support": "Original Weighted Support (%)",
    "original_weighted_provisional_grade": "Original Weighted Provisional Grade",
    "original_weighted_provisional_support": "Original Weighted Provisional Support (%)",
    "original_weighted_scores": "Original Weighted Class Support (JSON; fractions)",
    "final_combined_grade": "Final Combined Grade",
    "final_combined_support": "Final Combined Support (%)",
    "final_combined_scores": "Final Combined Class Support (JSON; fractions)",
    "score_scale": "Final Support Scale (raw weighted or normalized combined)",
    "geometric_class_c_points": "Geometric Class C Support (0-10 points; not percent)",
    "geometric_applied_score_shift": "Geometric Class C Applied Support Shift (fraction)",
    "width_changed_classification": "Width Changed Classification",
    "geometric_support": "Geometric Support Calculation (JSON)",
    "grading_configuration": "Per-fish Grading Configuration (JSON)",
    "part_detections": "Original Model 2 Part Detections (JSON)",
    "whole_fish_bbox": "Original Whole-fish Bounding Box (JSON)",
    "inspection_source": "Inspection Source and Frame (JSON)",
    "inspection_source_type": "Inspection Source Type",
    "source_frame_width_px": "Source Frame Width (px)",
    "source_frame_height_px": "Source Frame Height (px)",
    "complete_analysis": "Complete Fish Analysis (JSON)",
    "complete_fish_event": "Complete Structured Fish Event (JSON)",
}
for _grade, _prefix in (("Class A", "class_a"), ("Class B", "class_b"), ("Class C", "class_c"), ("Rejected", "rejected")):
    EXPORT_FIELDS[f"{_prefix}_original_weighted_score"] = f"{_grade} Original Weighted Support (%)"
    EXPORT_FIELDS[f"{_prefix}_final_combined_score"] = f"{_grade} Final Combined Support (%)"
for _region in ("body", "head", "tail", "whole_fish"):
    for _metric, _label in (("brown_ratio", "Brown Proxy"), ("dark_area_ratio", "Dark-area Proxy"),
                            ("discoloration_ratio", "Discoloration Proxy"), ("highlight_ratio", "Highlight Proxy"),
                            ("median_saturation", "Median Saturation"), ("median_value", "Median HSV Value")):
        EXPORT_FIELDS[f"{_region}_{_metric}"] = f"{_region.replace('_', ' ').title()} {_label} (%) (diagnostic only; rectangular ROI)"
    EXPORT_FIELDS.setdefault(f"{_region}_bbox_width_px", f"{_region.replace('_', ' ').title()} Bounding-box Width (px)")
    EXPORT_FIELDS.setdefault(f"{_region}_bbox_height_px", f"{_region.replace('_', ' ').title()} Bounding-box Height (px)")
    EXPORT_FIELDS[f"{_region}_measurement_rois"] = f"{_region.replace('_', ' ').title()} Measured ROIs and Frames (JSON)"
    EXPORT_FIELDS[f"{_region}_measurement_status"] = f"{_region.replace('_', ' ').title()} Image Measurement Status"
# These columns describe detector rectangles, not segmented fish or grading
# inputs. Keep the API field keys stable while making standalone CSV/XLSX
# headers carry the limitation even when other columns are omitted.
for _region in ("body", "head", "tail", "whole_fish"):
    for _metric in ("mean_hue", "median_hue", "mean_saturation", "mean_value", "yellow_ratio"):
        _field = f"{_region}_{_metric}"
        EXPORT_FIELDS[_field] += " (diagnostic only; rectangular ROI)"
DEFAULT_EXPORT_FIELDS = tuple(EXPORT_FIELDS)


def parse_timestamp(value: object) -> datetime | None:
    try:
        return datetime.fromisoformat(str(value))
    except (TypeError, ValueError):
        return None


def _number(value: object, *, scale: float = 1.0) -> float | str:
    try:
        number = float(value) * scale
    except (TypeError, ValueError):
        return ""
    return round(number, 4) if number == number and abs(number) != float("inf") else ""


def _part_values(event: Mapping[str, object], key: str) -> str:
    parts = event.get("parts")
    if not isinstance(parts, list):
        return ""
    return ", ".join(sorted({str(part.get(key, "")) for part in parts if isinstance(part, dict) and part.get(key)}))


def _analysis(event: Mapping[str, object]) -> dict[str, object]:
    value = event.get("analysis")
    return dict(value) if isinstance(value, dict) else {}


def _part_analysis(analysis: Mapping[str, object], region: str) -> dict[str, object]:
    parts = analysis.get("part_results")
    value = parts.get(region) if isinstance(parts, dict) else None
    return dict(value) if isinstance(value, dict) else {}


def _color_analysis(analysis: Mapping[str, object], region: str) -> dict[str, object]:
    colors = analysis.get("color")
    value = colors.get(region) if isinstance(colors, dict) else None
    return dict(value) if isinstance(value, dict) else {}


def _top_statistics(part: Mapping[str, object]) -> dict[str, object]:
    temporal = part.get("temporal_statistics")
    if not isinstance(temporal, dict):
        return {}
    top = temporal.get("top_grade_statistics")
    return dict(top) if isinstance(top, dict) else {}


def _traceability(analysis: Mapping[str, object]) -> dict[str, object]:
    trace = analysis.get("traceability")
    return dict(trace) if isinstance(trace, dict) else {}


def _json(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, allow_nan=False) if isinstance(value, (dict, list, tuple)) else ""


def _box_dimensions(value: object) -> tuple[object, object]:
    if not isinstance(value, (list, tuple)) or len(value) != 4:
        return None, None
    try:
        width, height = float(value[2]) - float(value[0]), float(value[3]) - float(value[1])
    except (TypeError, ValueError):
        return None, None
    return (width, height) if width > 0 and height > 0 else (None, None)


def _export_events(
    events: Iterable[dict[str, object]], *, source_metadata: Mapping[str, object] | None = None,
) -> list[dict[str, object]]:
    """Repair legacy measurement payloads at the structured export boundary.

    Keep the complete batch for geometric ownership/overlap checks. A saved
    width result is authoritative, including deliberately unavailable values;
    historical rows without image geometry are not assigned new measurements.
    """
    from src.inference.inspection_results import normalize_inspection_results, resolve_inspection_timestamp
    from src.inference.body_width import BodyWidthConfig

    values = list(events)
    candidates = []
    for event in values:
        analysis = _analysis(event)
        supporting = analysis.get("supporting_image_measurements")
        regions = supporting.get("regions") if isinstance(supporting, Mapping) else None
        has_geometry = _box_dimensions(event.get("bbox"))[0] is not None or isinstance(regions, Mapping) and bool(regions)
        width = analysis.get("body_width") if isinstance(analysis.get("body_width"), Mapping) else {}
        geometry = analysis.get("geometric_support")
        missing_stage = not isinstance(geometry, Mapping) or "original_grade" not in geometry or "final_grade" not in geometry
        recover_stage = False
        if missing_stage and width.get("threshold_configuration") and isinstance(analysis.get("weighted_scores"), Mapping):
            try:
                BodyWidthConfig.from_mapping(width["threshold_configuration"])
                recover_stage = True
            except (TypeError, ValueError):
                # Opaque historical configurations do not establish an
                # original pre-fusion decision; preserve their unknown fields.
                pass
        candidates.append((not width or not width.get("threshold_configuration") or recover_stage) and has_geometry)
    repaired = normalize_inspection_results(values, source_metadata=source_metadata) if any(candidates) else values
    result = []
    for original, normalized, candidate in zip(values, repaired, candidates):
        event = normalized if candidate else original
        if not event.get("timestamp"):
            timestamp = resolve_inspection_timestamp(event, source_metadata)
            if timestamp:
                event = {**event, "timestamp": timestamp}
        result.append(event)
    return result


def inspection_row(
    event: dict[str, object], *, source_metadata: Mapping[str, object] | None = None,
    _normalized: bool = False,
) -> dict[str, object]:
    """Map one completed fish to numeric, reconstructable export values."""

    if not _normalized:
        event = _export_events([event], source_metadata=source_metadata)[0]
    analysis = _analysis(event)
    geometry = analysis.get("geometric_support") if isinstance(analysis.get("geometric_support"), dict) else {}
    weighted = analysis.get("weighted_scores") if isinstance(analysis.get("weighted_scores"), dict) else {}
    final_grade = event.get("quality") or analysis.get("final_grade") or "Ungraded"
    best_grade = geometry.get("original_provisional_grade", analysis.get("best_evidence_class", analysis.get("provisional_grade")))
    best_score = geometry.get("original_provisional_support", analysis.get("best_evidence_score", analysis.get("provisional_score")))
    final_score = analysis.get("final_score")
    reason_codes = analysis.get("reason_codes")
    if isinstance(reason_codes, (list, tuple)):
        reason_codes_text = ", ".join(str(code) for code in reason_codes if code)
    else:
        reason_codes_text = str(analysis.get("verdict_reason_code") or "")
    model1 = analysis.get("model1_detection") if isinstance(analysis.get("model1_detection"), dict) else {}
    detection_confidence = model1.get("confidence")
    if detection_confidence is None:
        detection_confidence = event.get("final_confidence")
    model2_available = analysis.get("model2_available")
    trace = _traceability(analysis)
    temporal_stability = analysis.get("temporal_stability") if isinstance(analysis.get("temporal_stability"), dict) else {}
    track_stability = analysis.get("track_stability") if isinstance(analysis.get("track_stability"), dict) else {}
    frame_quality = analysis.get("frame_quality") if isinstance(analysis.get("frame_quality"), dict) else {}
    latest_quality = frame_quality.get("latest") if isinstance(frame_quality.get("latest"), dict) else {}
    best_frame = analysis.get("best_frame") if isinstance(analysis.get("best_frame"), dict) else {}
    damage = analysis.get("surface_damage") if isinstance(analysis.get("surface_damage"), dict) else {}
    conveyor_timing = analysis.get("conveyor_timing") if isinstance(analysis.get("conveyor_timing"), dict) else {}
    row: dict[str, object] = {
        "fish_id": event.get("fish_label", f"Fish #{event.get('track_id', event.get('id', ''))}"),
        "session_id": event.get("session_id", ""),
        "timestamp": event.get("timestamp", ""),
        "grade": final_grade,
        # Provisional evidence has its own column and must not be reported as
        # accepted final support when the fish was left Ungraded.
        "final_weighted_score": _number(final_score, scale=100),
        "best_evidence_class": best_grade or "",
        "best_evidence_score": _number(best_score, scale=100),
        "verdict_status": analysis.get("verdict_status", ""),
        "verdict_reason_code": analysis.get("verdict_reason_code", ""),
        "reason_codes": reason_codes_text,
        "verdict_reason": analysis.get("verdict_reason_text", ""),
        "second_grade": analysis.get("second_grade", ""),
        "second_support": _number(analysis.get("second_support"), scale=100),
        "grade_margin": _number(analysis.get("grade_margin"), scale=100),
        "manual_grade": event.get("manual_grade") or "",
        "manual_override": "Yes" if event.get("manual_override") else "No",
        "review_timestamp": event.get("review_timestamp") or "",
        "reviewer": event.get("reviewer") or "",
        "review_source": event.get("review_source") or "",
        "detection_confidence": _number(detection_confidence, scale=100),
        "surface_damage_status": damage.get("status", ""),
        "surface_damage_severity": damage.get("severity", ""),
        "surface_damage_confidence": _number(damage.get("confidence"), scale=100),
        "surface_damage_length_px": _number(damage.get("estimated_length_px")),
        "surface_damage_locations": json.dumps(damage.get("locations", []), ensure_ascii=False, sort_keys=True),
        "surface_damage_mask_source": damage.get("mask_source", ""),
        "grade_confidence": _number(event.get("quality_confidence"), scale=100),
        "observed_region_count": len(analysis.get("observed_regions", ())) if isinstance(analysis.get("observed_regions"), list) else "",
        "observed_regions": ", ".join(str(region) for region in analysis.get("observed_regions", ()) if region),
        "original_weight_coverage": _number(analysis.get("original_weight_coverage"), scale=100),
        "evidence_completeness": analysis.get("evidence_completeness", ""),
        "model2_observation_count": analysis.get("observation_count", ""),
        "model2_candidate_grading_frames": analysis.get("candidate_frame_count", temporal_stability.get("candidate_frame_count", "")),
        "model2_inference_interval_frames": trace.get("model2_inference_interval_frames", ""),
        "track_lifetime_detection_frames": track_stability.get("track_lifetime_detection_frames", track_stability.get("number_of_frames", "")),
        "model2_available": "Yes" if model2_available is True else "No" if model2_available is False else "",
        "sharpness_score": _number(latest_quality.get("sharpness_score") if isinstance(latest_quality, dict) else None),
        "best_frame_score": _number(best_frame.get("best_frame_score") if isinstance(best_frame, dict) else None),
        "best_frame_id": best_frame.get("best_frame_id", "") if isinstance(best_frame, dict) else "",
        "best_crop_path": best_frame.get("best_crop_path", "") if isinstance(best_frame, dict) else "",
        "class_a_weighted_score": _number(weighted.get("Class A") if isinstance(weighted, dict) else None, scale=100),
        "class_b_weighted_score": _number(weighted.get("Class B") if isinstance(weighted, dict) else None, scale=100),
        "class_c_weighted_score": _number(weighted.get("Class C") if isinstance(weighted, dict) else None, scale=100),
        "rejected_weighted_score": _number(weighted.get("Rejected") if isinstance(weighted, dict) else None, scale=100),
        "detected_regions": _part_values(event, "region"),
        "model2_classes": _part_values(event, "source_class_name"),
        "adjustment_applied": "Yes" if analysis.get("adjustments") else "No",
        "adjustment_reason": "; ".join(str(item.get("reason", "")) for item in analysis.get("adjustments", ()) if isinstance(item, dict)),
        "override_applied": "Yes" if isinstance(analysis.get("override"), dict) and analysis["override"].get("applied") else "No",
        "override_reason": analysis.get("override", {}).get("reason", "") if isinstance(analysis.get("override"), dict) else "",
        "model1_version": trace.get("model1_checkpoint", ""),
        "model1_hash": trace.get("model1_checkpoint_sha256", ""),
        "model2_version": trace.get("model2_checkpoint", ""),
        "model2_hash": trace.get("model2_checkpoint_sha256", ""),
        "grading_config_version": trace.get("grading_config_version", ""),
        "detection_threshold": _number(trace.get("detection_threshold")),
        "model2_detector_threshold": _number(trace.get("model2_detector_threshold", trace.get("quality_threshold"))),
        "part_evidence_threshold": _number(trace.get("minimum_part_evidence_threshold", trace.get("minimum_part_confidence"))),
        "quality_threshold": _number(trace.get("quality_threshold")),
        "final_verdict_threshold": _number(trace.get("final_verdict_threshold")),
        "minimum_coverage_setting": _number(trace.get("minimum_original_weight_coverage")),
        "minimum_grade_margin_setting": _number(trace.get("minimum_grade_margin")),
        "grading_mode": trace.get("grading_mode", ""),
        "weighting_method": analysis.get("weighting_method", ""),
        "tail_evaluation_status": analysis.get("tail_status", ""),
        "quality_inference_mode": trace.get("quality_inference_mode", ""),
        "roi_padding": _number(trace.get("roi_padding")),
        "tracker_backend": trace.get("tracker_backend", ""),
        "camera_settings_session_reference": event.get("session_id", ""),
        "final_explanation": " ".join(str(item) for item in analysis.get("explanation", ()) if item),
        "processing_time": event.get("processing_time_ms") if event.get("processing_time_ms") is not None else "",
        "decision_latency_ms": _number(conveyor_timing.get("first_detection_to_decision_ms")),
        "decision_deadline_seconds": _number(conveyor_timing.get("decision_deadline_seconds")),
        "decision_deadline_status": conveyor_timing.get("deadline_status") if isinstance(conveyor_timing.get("deadline_status"), str) else "",
        "decision_timing_scope": conveyor_timing.get("measurement_scope") if isinstance(conveyor_timing.get("measurement_scope"), str) else "",
        "supporting_image_measurements": json.dumps(analysis.get("supporting_image_measurements"), ensure_ascii=False, sort_keys=True, allow_nan=False) if isinstance(analysis.get("supporting_image_measurements"), dict) else "",
    }
    width = analysis.get("body_width") if isinstance(analysis.get("body_width"), dict) else {}
    width_configuration = width.get("threshold_configuration") if isinstance(width.get("threshold_configuration"), dict) else {}
    inspection_source = analysis.get("inspection_source") if isinstance(analysis.get("inspection_source"), dict) else {}
    original_scores = geometry.get("original_scores") if isinstance(geometry.get("original_scores"), dict) else {}
    final_scores = geometry.get("final_scores") if isinstance(geometry.get("final_scores"), dict) else weighted
    row.update({
        "original_weighted_grade": geometry.get("original_grade", ""),
        "original_weighted_support": _number(geometry.get("original_support"), scale=100),
        "original_weighted_provisional_grade": geometry.get("original_provisional_grade", ""),
        "original_weighted_provisional_support": _number(geometry.get("original_provisional_support"), scale=100),
        "original_weighted_scores": _json(original_scores) if original_scores else "",
        "final_combined_grade": geometry.get("final_grade", final_grade),
        "final_combined_support": _number(geometry.get("final_support", final_score), scale=100),
        "final_combined_scores": _json(final_scores),
        "score_scale": geometry.get("score_scale", "original_weighted_support" if weighted else "Unavailable"),
        "geometric_class_c_points": _number(geometry.get("class_c_points", width.get("class_c_supporting_points"))),
        "geometric_applied_score_shift": _number(geometry.get("applied_score_shift", width.get("applied_geometric_score_shift", width.get("applied_score_shift")))),
        "width_changed_classification": "Yes" if geometry.get("width_changed_classification") is True else "No" if geometry.get("width_changed_classification") is False else "Unavailable",
        "geometric_support": _json(geometry) if geometry else "",
        "grading_configuration": _json(analysis.get("grading_config")),
        "part_detections": _json(event.get("parts")),
        "whole_fish_bbox": _json(event.get("bbox", width.get("whole_fish_bbox"))),
        "complete_analysis": _json(analysis),
        "complete_fish_event": _json(event),
        "inspection_source": _json(inspection_source) if inspection_source else "",
        "inspection_source_type": inspection_source.get("type", ""),
        "source_frame_width_px": _number(inspection_source.get("frame_width_px")),
        "source_frame_height_px": _number(inspection_source.get("frame_height_px")),
        "body_width_scoring_method": geometry.get("scoring_method", width_configuration.get("scoring_method", "")),
        "body_width_relative_proxy_source": width.get("relative_width_proxy_source", width.get("measurement_method", "")),
        "body_width_configuration_source": width.get("threshold_configuration_source", "Recorded" if width_configuration else "Unavailable"),
    })
    for grade, prefix in (("Class A", "class_a"), ("Class B", "class_b"), ("Class C", "class_c"), ("Rejected", "rejected")):
        row[f"{prefix}_original_weighted_score"] = _number(original_scores.get(grade), scale=100)
        row[f"{prefix}_final_combined_score"] = _number(final_scores.get(grade), scale=100)
    for field, source in (
        ("whole_fish_bbox_width_px", "whole_fish_width_px"), ("whole_fish_bbox_height_px", "whole_fish_height_px"),
        ("body_bbox_width_px", "body_width_px"), ("body_bbox_height_px", "body_height_px"),
        ("relative_body_width_proxy", "relative_width_proxy"), ("body_width_initial_proxy", "initial_relative_width_proxy"),
        ("body_width_class_c_points", "class_c_supporting_points"), ("body_width_applied_body_score_shift", "applied_body_score_shift"),
        ("body_width_applied_geometric_score_shift", "applied_geometric_score_shift"),
        ("body_width_roi_proxy", "roi_relative_width_proxy"), ("body_width_detector_proxy", "bbox_relative_width_proxy"),
        ("whole_fish_roi_width_px", "whole_fish_roi_width_px"), ("whole_fish_roi_height_px", "whole_fish_roi_height_px"),
        ("body_roi_width_px", "body_roi_width_px"), ("body_roi_height_px", "body_roi_height_px"),
    ):
        row[field] = _number(width.get(source))
    points = row["body_width_class_c_points"]
    row["body_width_class_c_points_display"] = (
        f"{points:g} / 10" if isinstance(points, (int, float)) else str(width.get("score_status") or "Unavailable")
    )
    for field, source in (
        ("body_width_measurement_method", "measurement_method"), ("body_width_orientation", "orientation"),
        ("body_width_reliability", "reliability"), ("body_width_abnormality", "width_abnormality"),
        ("body_width_score_status", "score_status"), ("body_width_contribution_reason", "contribution_reason"),
        ("body_width_measurement_status", "measurement_status"), ("body_width_calibration_status", "calibration_status"),
    ):
        row[field] = width.get(source) or "Unavailable"
    reasons = width.get("reliability_reasons")
    row["body_width_reliability_reasons"] = "; ".join(str(reason) for reason in reasons) if isinstance(reasons, (list, tuple)) else ""
    for field, source in (("body_width_scoring_enabled", "scoring_enabled"), ("body_width_score_contributed", "score_contributed"),
                          ("body_width_score_applied", "score_applied"), ("body_width_measurement_available", "measurement_available"),
                          ("body_width_reliable_for_scoring", "reliable_for_scoring")):
        row[field] = "Yes" if width.get(source) is True else "No" if width.get(source) is False else "Unavailable"
    for field, source, expected in (
        ("body_width_normal_reference_range", "normal_reference_range", (list, tuple)),
        ("body_width_threshold_configuration", "threshold_configuration", (dict,)),
    ):
        value = width.get(source)
        row[field] = json.dumps(value, ensure_ascii=False, sort_keys=True, allow_nan=False) if isinstance(value, expected) else ""
    row["body_width_analysis"] = json.dumps(width, ensure_ascii=False, sort_keys=True, allow_nan=False) if width else ""
    fish_geometry = analysis.get("fish_geometry") if isinstance(analysis.get("fish_geometry"), dict) else {}
    physical = fish_geometry.get("physical_calibration") if isinstance(fish_geometry.get("physical_calibration"), dict) else {}
    # Withheld silhouette candidates and historical bounding boxes are never
    # substituted for actual per-fish measurements or physical conversions.
    measured = fish_geometry.get("status") == "measured"
    physical_applied = measured and physical.get("validated") is True and physical.get("status") == "Applied"
    geometry_reasons = fish_geometry.get("reasons")
    row.update({
        "fish_length_px": _number(fish_geometry.get("fish_length_px")) if measured else "",
        "max_body_width_px": _number(fish_geometry.get("max_body_width_px")) if measured else "",
        "relative_fish_body_width_ratio": _number(fish_geometry.get("relative_width_ratio")) if measured else "",
        "fish_geometry_status": fish_geometry.get("status") or "unavailable",
        "fish_geometry_reliability": fish_geometry.get("reliability") or "Unavailable",
        "fish_geometry_reasons": "; ".join(str(reason) for reason in geometry_reasons) if isinstance(geometry_reasons, (list, tuple)) else "No per-fish silhouette measurement recorded",
        "fish_geometry_method": fish_geometry.get("method") or "Unavailable",
        "fish_geometry_orientation_degrees": _number(fish_geometry.get("orientation_degrees")) if measured else "",
        "fish_geometry_measurement_validated": "Yes" if fish_geometry.get("measurement_validated") is True else "Not validated",
        "fish_geometry_used_for_grading": "Yes" if fish_geometry.get("used_for_grading") is True else "No",
        "fish_length_cm": _number(fish_geometry.get("fish_length_cm")) if physical_applied else "",
        "max_body_width_cm": _number(fish_geometry.get("max_body_width_cm")) if physical_applied else "",
        "physical_calibration_status": physical.get("status") or "Not calibrated",
        "physical_calibration": _json(physical) if physical else "",
        "fish_geometry_analysis": _json(fish_geometry) if fish_geometry else "",
    })
    winning_grade = str(final_grade if final_grade != "Ungraded" else (best_grade or ""))
    for region, prefix in (("Body", "body"), ("Head", "head"), ("Tail", "tail")):
        part, color, stats = _part_analysis(analysis, region), _color_analysis(analysis, region), _top_statistics(_part_analysis(analysis, region))
        contributions = part.get("contributions") if isinstance(part.get("contributions"), dict) else {}
        grade_evidence = part.get("grade_evidence") if isinstance(part.get("grade_evidence"), dict) else {}
        present = part.get("present")
        row[f"{prefix}_observed"] = "Yes" if present is True else "No" if present is False else ""
        row[f"{prefix}_presence"] = part.get("visibility_state", part.get("status", part.get("observation_status", part.get("missing_status", "unknown"))))
        row[f"{prefix}_presence_confidence"] = _number(part.get("presence_confidence"), scale=100)
        row[f"{prefix}_class_a_evidence"] = _number(grade_evidence.get("Class A"), scale=100)
        row[f"{prefix}_class_b_evidence"] = _number(grade_evidence.get("Class B"), scale=100)
        row[f"{prefix}_class_c_evidence"] = _number(grade_evidence.get("Class C"), scale=100)
        row[f"{prefix}_rejected_evidence"] = _number(grade_evidence.get("Rejected"), scale=100)
        row[f"{prefix}_grade"] = part.get("grade", "")
        row[f"{prefix}_grade_confidence"] = _number(part.get("grade_confidence"), scale=100)
        row[f"{prefix}_weight"] = _number(part.get("original_weight", part.get("weight")))
        row[f"{prefix}_effective_weight"] = _number(part.get("effective_weight", part.get("normalized_weight")))
        row[f"{prefix}_weighted_contribution"] = _number(contributions.get(winning_grade) if isinstance(contributions, dict) else None, scale=100)
        row[f"{prefix}_class_a_contribution"] = _number(contributions.get("Class A") if isinstance(contributions, dict) else None, scale=100)
        row[f"{prefix}_class_b_contribution"] = _number(contributions.get("Class B") if isinstance(contributions, dict) else None, scale=100)
        row[f"{prefix}_class_c_contribution"] = _number(contributions.get("Class C") if isinstance(contributions, dict) else None, scale=100)
        row[f"{prefix}_rejected_contribution"] = _number(contributions.get("Rejected") if isinstance(contributions, dict) else None, scale=100)
        row[f"{prefix}_observation_count"] = stats.get("number_of_valid_observations", "")
        row[f"{prefix}_mean_confidence"] = _number(stats.get("mean_confidence"), scale=100)
        row[f"{prefix}_max_confidence"] = _number(stats.get("max_confidence"), scale=100)
        row[f"{prefix}_std_dev"] = _number(stats.get("standard_deviation"), scale=100)
        row[f"{prefix}_mean_hue"] = _number(color.get("mean_hue_deg"))
        row[f"{prefix}_median_hue"] = _number(color.get("median_hue_deg"))
        row[f"{prefix}_mean_saturation"] = _number(color.get("mean_saturation"), scale=100)
        row[f"{prefix}_mean_value"] = _number(color.get("mean_value"), scale=100)
        row[f"{prefix}_yellow_ratio"] = _number(color.get("yellow_ratio_proxy"), scale=100)
    whole = _color_analysis(analysis, "Whole Fish")
    row["whole_fish_mean_hue"] = _number(whole.get("mean_hue_deg"))
    row["whole_fish_median_hue"] = _number(whole.get("median_hue_deg"))
    row["whole_fish_mean_saturation"] = _number(whole.get("mean_saturation"), scale=100)
    row["whole_fish_mean_value"] = _number(whole.get("mean_value"), scale=100)
    row["whole_fish_yellow_ratio"] = _number(whole.get("yellow_ratio_proxy"), scale=100)
    supporting = analysis.get("supporting_image_measurements")
    regions = supporting.get("regions", {}) if isinstance(supporting, dict) else {}
    detections = event.get("parts", [])
    for region, prefix in (("Body", "body"), ("Head", "head"), ("Tail", "tail"), ("Whole Fish", "whole_fish")):
        measured = regions.get(region, {}) if isinstance(regions, dict) else {}
        measured = measured if isinstance(measured, dict) else {}
        color = _color_analysis(analysis, region)
        values = measured.get("values") if isinstance(measured.get("values"), dict) else color
        # Diagnostic-only snapshots can exist even when the grading policy
        # rejects the frame. Reuse their values rather than hiding valid data.
        for metric, source, scale in (("mean_hue", "mean_hue_deg", 1), ("median_hue", "median_hue_deg", 1),
                                      ("mean_saturation", "mean_saturation", 100), ("mean_value", "mean_value", 100),
                                      ("yellow_ratio", "yellow_ratio_proxy", 100)):
            row[f"{prefix}_{metric}"] = _number(values.get(source), scale=scale)
        rois = measured.get("frame_rois")
        row[f"{prefix}_measurement_rois"] = _json(rois)
        row[f"{prefix}_measurement_status"] = measured.get("status", color.get("status", "Unavailable"))
        for metric, source in (("brown_ratio", "brown_ratio_proxy"), ("dark_area_ratio", "dark_patch_ratio_proxy"),
                               ("discoloration_ratio", "discoloration_coverage_proxy"), ("highlight_ratio", "specular_ratio_proxy"),
                               ("median_saturation", "median_saturation"), ("median_value", "median_value")):
            row[f"{prefix}_{metric}"] = _number(values.get(source), scale=100)
        if row.get(f"{prefix}_bbox_width_px", "") == "":
            if region == "Whole Fish":
                bbox = event.get("bbox")
            else:
                candidates = [part for part in detections if isinstance(part, dict) and part.get("region") == region] if isinstance(detections, list) else []
                # Select the same strongest detector rectangle used by the image diagnostics.
                selected = max(candidates, key=lambda part: float(part.get("confidence") or 0), default={})
                bbox = selected.get("bbox")
            box_width, box_height = _box_dimensions(bbox)
            row[f"{prefix}_bbox_width_px"] = _number(box_width)
            row[f"{prefix}_bbox_height_px"] = _number(box_height)
    return row


def filter_events(events: Iterable[dict[str, object]], *, range_name: str = "current_session", start_date: str | None = None, end_date: str | None = None) -> list[dict[str, object]]:
    values = list(events)
    if range_name == "current_session":
        return values
    if range_name == "today":
        today = date.today()
        return [event for event in values if (parsed := parse_timestamp(event.get("timestamp"))) and parsed.date() == today]
    if range_name != "custom":
        raise ValueError("Unsupported export range.")
    try:
        lower = date.fromisoformat(start_date) if start_date else None
        upper = date.fromisoformat(end_date) if end_date else None
    except ValueError as exc:
        raise ValueError("Custom export dates must use YYYY-MM-DD.") from exc
    if lower and upper and lower > upper:
        raise ValueError("The start date must not be after the end date.")
    return [
        event for event in values
        if (parsed := parse_timestamp(event.get("timestamp"))) is not None
        and (not lower or parsed.date() >= lower) and (not upper or parsed.date() <= upper)
    ]


def _csv_spreadsheet_text(value: object) -> object:
    """Keep user-supplied text from becoming a formula in spreadsheet apps."""

    if isinstance(value, str) and value and (
        value[0] in "\t\r\n" or value.lstrip().startswith(("=", "+", "-", "@"))
    ):
        return "'" + value
    return value


def make_csv(
    events: Iterable[dict[str, object]], fields: Iterable[str], *,
    source_metadata: Mapping[str, object] | None = None,
) -> bytes:
    selected = tuple(fields)
    buffer = StringIO(newline="")
    writer = csv.DictWriter(buffer, fieldnames=[EXPORT_FIELDS[field] for field in selected], extrasaction="ignore")
    writer.writeheader()
    for event in _export_events(events, source_metadata=source_metadata):
        row = inspection_row(event, _normalized=True)
        writer.writerow({EXPORT_FIELDS[field]: _csv_spreadsheet_text(row[field]) for field in selected})
    return buffer.getvalue().encode("utf-8-sig")


def _style_header(sheet) -> None:
    from openpyxl.styles import Font, PatternFill
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")


def make_xlsx(
    events: Iterable[dict[str, object]],
    fields: Iterable[str],
    summary: dict[str, object],
    grading_rules: Mapping[str, object] | None = None,
    *,
    session_configuration: Mapping[str, object] | None = None,
    session_configurations: Iterable[Mapping[str, object]] | None = None,
    source_metadata: Mapping[str, object] | None = None,
) -> bytes:
    """Create explainable inspection sheets plus an optional durable session snapshot."""

    try:
        from openpyxl import Workbook
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("Excel export requires openpyxl. Install the project requirements.") from exc
    selected, event_rows = tuple(fields), _export_events(events, source_metadata=source_metadata)
    workbook = Workbook()
    inspections = workbook.active
    inspections.title = "Fish Inspections"
    inspections.append([EXPORT_FIELDS[field] for field in selected])
    _style_header(inspections)
    for event in event_rows:
        row = inspection_row(event, _normalized=True)
        inspections.append([row[field] for field in selected])
    inspections.freeze_panes = "A2"
    for column in inspections.columns:
        inspections.column_dimensions[column[0].column_letter].width = min(48, max(14, max(len(str(cell.value or "")) for cell in column) + 2))

    summary_sheet = workbook.create_sheet("Session Summary")
    summary_sheet.append(["Metric", "Value"])
    _style_header(summary_sheet)
    for label, value in summary.items():
        summary_sheet.append([label, value])
    summary_sheet.column_dimensions["A"].width, summary_sheet.column_dimensions["B"].width = 42, 26

    feature_sheet = workbook.create_sheet("Feature Summary")
    feature_sheet.append(["Feature", "Count", "Percentage of fish", "Source"])
    _style_header(feature_sheet)
    total = len(event_rows)
    for region in ("Head", "Body", "Tail"):
        count = sum(int(_part_analysis(_analysis(event), region).get("present") is True) for event in event_rows)
        feature_sheet.append([f"{region} confidently observed", count, round(count / total * 100, 2) if total else 0.0, "Model 2 part detector"])
    measured_color = sum(
        1 for event in event_rows
        if any(_color_analysis(_analysis(event), region).get("status") == "measured"
               for region in ("Whole Fish", "Body", "Head", "Tail"))
    )
    feature_sheet.append(["HSV colour metrics (diagnostic only)", measured_color,
                          round(measured_color / total * 100, 2) if total else 0.0,
                          "Rectangular ROI proxy; not used for grading"])
    feature_sheet.append(["Trained defect labels", 0, 0.0, "None in supplied Model 1/Model 2 checkpoints"])
    for column, width in {"A": 30, "B": 14, "C": 22, "D": 56}.items():
        feature_sheet.column_dimensions[column].width = width

    configurations = list(session_configurations) if session_configurations is not None else None
    rules = dict(grading_rules or {})
    weights = rules.get("part_weights") if isinstance(rules.get("part_weights"), dict) else {}
    rule_sheet = workbook.create_sheet("Grade Calculation Rules")
    rule_sheet.append(["Rule", "Value", "Status"])
    _style_header(rule_sheet)
    if configurations is not None:
        rule_sheet.append(["Configuration provenance", "Multiple sessions; see Session Configurations and Row Session Mapping.", "Per session"])
    elif not rules:
        rule_sheet.append(["Configuration provenance", "The session grading configuration was not recorded.", "Unavailable"])
    else:
        for region in ("Body", "Head", "Tail"):
            rule_sheet.append([f"{region} original weight", weights.get(region, ""), "Recorded" if region in weights else "Not recorded"])
        rule_keys = (
            ("Configuration version", "version"), ("Aggregation", "aggregation"), ("Unavailable region", "unavailable_region"),
            ("Body required", "require_body"), ("Grading mode", "grading_mode"), ("Minimum regions", "minimum_regions_observed"),
            ("Minimum original coverage", "minimum_original_weight_coverage"), ("Minimum valid frames", "minimum_track_observations"),
            ("Rejected override threshold", "rejected_override_threshold"), ("HSV adjustments", "color_adjustments_enabled"),
            ("Defect adjustments", "defect_adjustments_enabled"),
        )
        for label, key in rule_keys:
            value = rules.get(key)
            rule_sheet.append([label, "" if value is None else value, "Recorded" if key in rules else "Not recorded"])
        threshold_key = "active_final_verdict_threshold" if "active_final_verdict_threshold" in rules else "final_verdict_threshold"
        rule_sheet.append(["Final verdict threshold", rules.get(threshold_key, ""), "Recorded" if threshold_key in rules else "Not recorded"])
        width_rules = rules.get("body_width")
        if isinstance(width_rules, Mapping):
            rule_sheet.append(["Geometric scoring method", width_rules.get("scoring_method", "body_component_v1"), "Recorded"])
            rule_sheet.append(["Width scoring enabled", width_rules.get("scoring_enabled", False), "Recorded"])
            rule_sheet.append(["Width calibration validated", width_rules.get("calibration_validated", False), "Recorded"])
            rule_sheet.append(["Complete width configuration", _json(dict(width_rules)), "Recorded"])
    rule_sheet.column_dimensions["A"].width, rule_sheet.column_dimensions["B"].width, rule_sheet.column_dimensions["C"].width = 30, 96, 42

    review_sheet = workbook.create_sheet("Review Queue")
    review_headers = ["Fish ID", "Timestamp", "Best Evidence Class", "Best Score (%)", "Evidence Coverage (%)", "Observed Regions", "Review Reason", "Manual Grade"]
    review_sheet.append(review_headers)
    _style_header(review_sheet)
    for event in event_rows:
        analysis = _analysis(event)
        if analysis.get("verdict_status") != "NEEDS_REVIEW" or event.get("manual_override"):
            continue
        row = inspection_row(event, _normalized=True)
        review_sheet.append([
            row["fish_id"], row["timestamp"], row["best_evidence_class"], row["best_evidence_score"],
            row["original_weight_coverage"], row["observed_regions"], row["verdict_reason"], row["manual_grade"],
        ])
    for column, width in {"A": 14, "B": 27, "C": 22, "D": 16, "E": 22, "F": 28, "G": 74, "H": 18}.items():
        review_sheet.column_dimensions[column].width = width
    review_sheet.freeze_panes = "A2"

    if session_configuration is not None:
        configuration_sheet = workbook.create_sheet("Session Configuration")
        configuration_sheet.append(["Field", "Value"])
        _style_header(configuration_sheet)
        for key, value in session_configuration.items():
            if isinstance(value, (dict, list, tuple)):
                rendered = json.dumps(value, ensure_ascii=False, sort_keys=True, default=str)
            else:
                rendered = "" if value is None else value
            configuration_sheet.append([str(key), rendered])
        configuration_sheet.column_dimensions["A"].width = 32
        configuration_sheet.column_dimensions["B"].width = 100
        configuration_sheet.freeze_panes = "A2"
    if configurations is not None:
        configuration_sheet = workbook.create_sheet("Session Configurations")
        configuration_sheet.append(["Session ID", "Model 1 SHA256", "Model 2 SHA256", "Grading Config Hash", "Grading Configuration", "Status"])
        _style_header(configuration_sheet)
        for configuration in configurations:
            policy = configuration.get("grading_config")
            configuration_sheet.append([
                configuration.get("session_id") or "Not recorded", configuration.get("model1_sha256") or "",
                configuration.get("model2_sha256") or "", configuration.get("grading_config_hash") or "",
                json.dumps(policy, ensure_ascii=False, sort_keys=True, default=str) if isinstance(policy, Mapping) else "",
                "Recorded" if isinstance(policy, Mapping) and policy else "Unavailable",
            ])
        for column, width in {"A": 36, "B": 66, "C": 66, "D": 66, "E": 100, "F": 20}.items():
            configuration_sheet.column_dimensions[column].width = width
        configuration_sheet.freeze_panes = "A2"
        mapping_sheet = workbook.create_sheet("Row Session Mapping")
        mapping_sheet.append(["Fish Inspections row", "Session ID"])
        _style_header(mapping_sheet)
        for row_number, event in enumerate(event_rows, start=2):
            mapping_sheet.append([row_number, event.get("session_id") or "Not recorded"])
        mapping_sheet.column_dimensions["A"].width, mapping_sheet.column_dimensions["B"].width = 25, 40
        mapping_sheet.freeze_panes = "A2"
    # Workbook text originates partly in operator-entered/session data. Force
    # formula-looking strings to remain literal values in every sheet.
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    cell.data_type = "s"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
