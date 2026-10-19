"""Guard the distinction between grading evidence and image diagnostics."""

import csv
from io import BytesIO, StringIO
from pathlib import Path
import unittest

import numpy as np
from openpyxl import load_workbook

from src.api.export import EXPORT_FIELDS, make_csv, make_xlsx
from src.api.domain import FEATURE_CAPABILITIES
from src.inference.grading_engine import GradingConfig, WeightedGradingEngine
from src.inference.part_types import PartDetection


ROOT = Path(__file__).resolve().parents[1]


class FeatureClaimsTests(unittest.TestCase):
    def test_diagnostic_headers_are_standalone_and_field_keys_stable(self):
        fields = ("body_mean_value", "whole_fish_yellow_ratio", "grade_confidence")
        header = next(csv.reader(StringIO(make_csv([], fields).decode("utf-8-sig"))))
        self.assertEqual(header, [EXPORT_FIELDS[field] for field in fields])
        for title in header[:2]:
            self.assertIn("diagnostic only", title)
            self.assertIn("rectangular ROI", title)
        self.assertIn("Support", header[2])

    def test_feature_summary_counts_only_measured_fish(self):
        events = [
            {"analysis": {"color": {"Body": {"status": "measured", "mean_value": 0.5}}}},
            {"analysis": {"color": {"Body": {"status": "unavailable"}}}},
            {"analysis": {}},
        ]
        workbook = load_workbook(BytesIO(make_xlsx(events, ("fish_id",), {})), data_only=True)
        rows = list(workbook["Feature Summary"].values)
        color = next(row for row in rows if row[0].startswith("HSV colour metrics"))
        self.assertEqual(color[1:3], (1, 33.33))
        self.assertIn("not used for grading", color[3])
        empty = load_workbook(BytesIO(make_xlsx([], ("fish_id",), {})), data_only=True)
        self.assertEqual(list(empty["Feature Summary"].values)[4][1:3], (0, 0))

    def test_adjustment_switches_do_not_create_a_grading_rule(self):
        image = np.full((32, 32, 3), 160, dtype=np.uint8)
        parts = [PartDetection.from_source_class(0, 0.9, (0, 0, 32, 32))]
        baseline = WeightedGradingEngine(GradingConfig()).evaluate(1, image, parts, stabilize=False)
        switches = WeightedGradingEngine(GradingConfig(color_adjustments_enabled=True, defect_adjustments_enabled=True)).evaluate(1, image, parts, stabilize=False)
        self.assertEqual(baseline.weighted_scores, switches.weighted_scores)
        self.assertEqual(baseline.final_grade, switches.final_grade)
        self.assertEqual(switches.adjustments, ())
        self.assertTrue(any("scores remain unchanged" in line for line in switches.explanation))

    def test_interface_names_proxies_and_never_advertises_defect_capabilities(self):
        html = (ROOT / "frontend/index.html").read_text(encoding="utf-8")
        javascript = (ROOT / "frontend/app.js").read_text(encoding="utf-8")
        self.assertIn("the current models do not detect belly defects", html)
        self.assertIn("Diagnostic only", html)
        self.assertIn("do not affect grading", html)
        self.assertIn("calibrated probability of a correct grade", html)
        self.assertIn("not validated defect or shininess detection", javascript)
        self.assertIn("HSV value (brightness proxy)", javascript)
        self.assertEqual(set(FEATURE_CAPABILITIES), {"Head", "Body", "Tail"})


if __name__ == "__main__":
    unittest.main()
