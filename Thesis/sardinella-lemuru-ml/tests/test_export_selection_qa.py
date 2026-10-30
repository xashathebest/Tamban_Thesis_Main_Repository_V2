"""Exported inspection columns follow the operator's selected order."""
import csv
from io import BytesIO, StringIO
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from src.api import app
from src.api.export import DEFAULT_EXPORT_FIELDS, EXPORT_FIELDS, inspection_row, make_csv, make_xlsx


@pytest.fixture
def client(monkeypatch):
    monkeypatch.setattr(app, "_durable_events", lambda **_: [{"track_id": 7, "timestamp": "2026-01-01T00:00:00+00:00", "quality": "Class A", "analysis": {}}])
    monkeypatch.setattr(app, "_status_payload", lambda: {"grading_rules": {}})
    monkeypatch.setattr(app, "inspection_sessions", SimpleNamespace(session_id=None))
    monkeypatch.setattr(app, "inspection_database", SimpleNamespace(session=lambda _: None))
    return TestClient(app.app)


def headers(response, export_format):
    if export_format == "csv":
        return next(csv.reader(StringIO(response.content.decode("utf-8-sig"))))
    workbook = load_workbook(BytesIO(response.content), data_only=True)
    assert {"Fish Inspections", "Session Summary", "Feature Summary", "Grade Calculation Rules", "Review Queue"} <= set(workbook.sheetnames)
    return [cell.value for cell in workbook["Fish Inspections"][1]]


@pytest.mark.parametrize("export_format", ["csv", "xlsx"])
def test_selected_fields_are_the_only_inspection_columns_in_requested_order(client, export_format):
    fields = ["grade", "fish_id", "timestamp"]
    response = client.post("/api/export", json={"format": export_format, "fields": fields})
    assert response.status_code == 200
    assert headers(response, export_format) == [EXPORT_FIELDS[field] for field in fields]


@pytest.mark.parametrize("export_format", ["csv", "xlsx"])
def test_explicit_full_evidence_export_works_with_no_checked_fields(client, export_format):
    response = client.post("/api/export", json={"format": export_format, "fields": [], "complete_analysis": True})
    assert response.status_code == 200
    assert headers(response, export_format) == [EXPORT_FIELDS[field] for field in DEFAULT_EXPORT_FIELDS]


@pytest.mark.parametrize("payload", [
    {},
    {"fields": []},
    {"fields": ["fish_id", "fish_id"]},
    {"fields": ["fish_id"], "complete_analysis": "true"},
])
def test_invalid_selection_is_rejected(client, payload):
    assert client.post("/api/export", json=payload).status_code == 422


def test_exported_user_text_remains_text_in_spreadsheets():
    event = {"track_id": 7, "fish_label": 'Fish, "Mañana"\nBatch', "reviewer": '=1+1',
             "timestamp": "2026-10-18T12:00:00-04:00", "quality": "Class A", "analysis": {}}
    fields = ["fish_id", "timestamp", "reviewer"]
    csv_rows = list(csv.reader(StringIO(make_csv([event], fields).decode("utf-8-sig"))))
    assert csv_rows[1] == [event["fish_label"], event["timestamp"], "'=1+1"]
    workbook = load_workbook(BytesIO(make_xlsx([event], fields, {}, session_configuration={"note": "=2+2"})), data_only=False)
    row = workbook["Fish Inspections"][2]
    assert [cell.value for cell in row] == [event["fish_label"], event["timestamp"], event["reviewer"]]
    assert row[2].data_type == "s"
    assert workbook["Session Configuration"]["B2"].value == "=2+2"
    assert workbook["Session Configuration"]["B2"].data_type == "s"


def test_empty_exports_keep_headers():
    assert len(list(csv.reader(StringIO(make_csv([], ["fish_id"]).decode("utf-8-sig"))))) == 1
    workbook = load_workbook(BytesIO(make_xlsx([], ["fish_id"], {})), data_only=True)
    assert list(workbook["Fish Inspections"].values) == [(EXPORT_FIELDS["fish_id"],)]


def test_export_preserves_available_detection_confidence_without_inventing_model2_status():
    event = {"final_confidence": 0.91, "analysis": {}}
    row = inspection_row(event)
    assert row["detection_confidence"] == 91.0
    assert row["model2_available"] == ""
    event["analysis"] = {"model1_detection": {"confidence": 0.75}, "model2_available": False}
    row = inspection_row(event)
    assert row["detection_confidence"] == 75.0
    assert row["model2_available"] == "No"
    event["analysis"]["model2_available"] = True
    assert inspection_row(event)["model2_available"] == "Yes"


def test_export_includes_stable_conveyor_decision_timing_without_inventing_missing_values():
    event = {
        "analysis": {
            "conveyor_timing": {
                "first_detection_to_decision_ms": 482.125,
                "decision_deadline_seconds": 0.75,
                "deadline_status": "met",
                "measurement_scope": "First persistent Model 1 detection to virtual count-line decision.",
            }
        }
    }
    row = inspection_row(event)
    assert row["decision_latency_ms"] == 482.125
    assert row["decision_deadline_seconds"] == 0.75
    assert row["decision_deadline_status"] == "met"
    assert row["decision_timing_scope"] == "First persistent Model 1 detection to virtual count-line decision."
    assert all(field in DEFAULT_EXPORT_FIELDS for field in (
        "decision_latency_ms", "decision_deadline_seconds", "decision_deadline_status", "decision_timing_scope",
    ))
    csv_rows = list(csv.reader(StringIO(make_csv([event], ["decision_latency_ms", "decision_deadline_status"]).decode("utf-8-sig"))))
    assert csv_rows == [
        [EXPORT_FIELDS["decision_latency_ms"], EXPORT_FIELDS["decision_deadline_status"]],
        ["482.125", "met"],
    ]
    assert inspection_row({"analysis": {}})["decision_latency_ms"] == ""
    assert inspection_row({"analysis": {}})["decision_deadline_status"] == ""


def test_ungraded_export_keeps_best_evidence_separate_from_final_support():
    event = {"quality": "Ungraded", "analysis": {"final_score": None, "provisional_score": 0.64,
             "provisional_grade": "Class A", "part_results": {"Head": {"present": False}}}}
    row = inspection_row(event)
    assert row["final_weighted_score"] == ""
    assert row["best_evidence_score"] == 64.0
    assert row["head_observed"] == "No"
    assert row["body_observed"] == ""


def test_export_review_queue_excludes_already_reviewed_results():
    pending = {"track_id": 1, "analysis": {"verdict_status": "NEEDS_REVIEW"}}
    reviewed = {"track_id": 2, "manual_override": True, "manual_grade": "Class A",
                "analysis": {"verdict_status": "NEEDS_REVIEW"}}
    workbook = load_workbook(BytesIO(make_xlsx([pending, reviewed], ["fish_id", "manual_grade"], {})), data_only=True)
    assert [row[0] for row in workbook["Review Queue"].iter_rows(min_row=2, values_only=True)] == ["Fish #1"]
    # The reviewed fish is retained separately in the exported inspection rows.
    assert list(workbook["Fish Inspections"].values)[2] == ("Fish #2", "Class A")


@pytest.mark.parametrize("value", ["=1+1", "  =1+1", "+1", "-1", "@name", "\ttext"])
def test_csv_neutralizes_spreadsheet_formula_prefixes(value):
    rows = list(csv.reader(StringIO(make_csv([{"reviewer": value}], ["reviewer"]).decode("utf-8-sig"))))
    assert rows[1] == ["'" + value]
