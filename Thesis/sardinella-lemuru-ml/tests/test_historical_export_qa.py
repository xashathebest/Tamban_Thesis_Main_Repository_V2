"""Historic reports must retain their own session policy and model identity."""
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from src.api import app
from src.api.database import InspectionDatabase, InspectionSessionContext


@pytest.fixture
def history(tmp_path, monkeypatch):
    db = InspectionDatabase(tmp_path / "history.sqlite3")
    old = db.create_session({"started_at": "2020-01-01T00:00:00+00:00", "model2_sha256": "old-hash",
                             "grading_config": {"part_weights": {"Body": 0.7, "Head": 0.2, "Tail": 0.1}, "version": "old", "final_verdict_threshold": 0.3}})
    sessions = InspectionSessionContext(db)
    new = sessions.start({"started_at": "2021-01-01T00:00:00+00:00", "model2_sha256": "new-hash",
                          "grading_config": {"part_weights": {"Body": 0.5, "Head": 0.3, "Tail": 0.2}, "version": "new", "final_verdict_threshold": 0.8}})
    for sid, track in ((old, 1), (new, 2)):
        db.upsert_event(sid, {"track_id": track, "timestamp": "2021-01-01T00:00:00+00:00", "quality": "Class A", "analysis": {}})
    monkeypatch.setattr(app, "inspection_database", db)
    monkeypatch.setattr(app, "inspection_sessions", sessions)
    monkeypatch.setattr(app, "_status_payload", lambda: {"grading_rules": {"part_weights": {"Body": 0.5}, "final_verdict_threshold": 0.8}})
    return TestClient(app.app), db, old, new


def sheet_values(sheet):
    return {row[0]: row[1] for row in sheet.iter_rows(min_row=2, values_only=True) if row[0] is not None}


def test_historical_xlsx_uses_its_own_rules_and_hash(history):
    client, _, old, _ = history
    response = client.post("/api/export", json={"format": "xlsx", "range": "current_session", "session_id": old, "fields": ["fish_id"]})
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), data_only=True)
    rules = sheet_values(workbook["Grade Calculation Rules"])
    assert rules["Body original weight"] == 0.7
    assert rules["Final verdict threshold"] == 0.3
    assert sheet_values(workbook["Session Configuration"])["model2_sha256"] == "old-hash"


def test_mixed_xlsx_has_per_session_policy_and_row_mapping(history):
    client, _, old, new = history
    response = client.post("/api/export", json={"format": "xlsx", "range": "all", "fields": ["fish_id"]})
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), data_only=True)
    assert workbook["Fish Inspections"].max_column == 1
    assert {"Session Configurations", "Row Session Mapping"} <= set(workbook.sheetnames)
    records = list(workbook["Session Configurations"].iter_rows(min_row=2, values_only=True))
    assert {row[0] for row in records} == {old, new}
    assert any("old-hash" in str(row) and '"Body": 0.7' in str(row) for row in records)
    assert any("new-hash" in str(row) and '"Body": 0.5' in str(row) for row in records)
    mapping = list(workbook["Row Session Mapping"].iter_rows(min_row=2, values_only=True))
    assert {row[1] for row in mapping} == {old, new}
    assert {workbook["Fish Inspections"].cell(row[0], 1).value: row[1] for row in mapping} == {"Fish #1": old, "Fish #2": new}
    assert sheet_values(workbook["Grade Calculation Rules"])["Configuration provenance"].startswith("Multiple sessions")


def test_mixed_csv_requires_xlsx_for_policy_provenance(history):
    client, _, _, _ = history
    response = client.post("/api/export", json={"format": "csv", "range": "all", "fields": ["fish_id"]})
    assert response.status_code == 422 and "Excel" in response.json()["detail"]


def test_unknown_session_returns_404(history):
    client, _, _, _ = history
    response = client.post("/api/export", json={"format": "xlsx", "range": "current_session", "session_id": "missing", "fields": ["fish_id"]})
    assert response.status_code == 404


def test_missing_policy_is_reported_without_inventing_defaults(history):
    client, db, _, _ = history
    unknown = db.create_session({"started_at": "2019-01-01T00:00:00+00:00"})
    response = client.post("/api/export", json={"format": "xlsx", "session_id": unknown, "fields": ["fish_id"]})
    workbook = load_workbook(BytesIO(response.content), data_only=True)
    rules = sheet_values(workbook["Grade Calculation Rules"])
    assert "not recorded" in rules["Configuration provenance"]
    assert "Body original weight" not in rules


def test_session_policy_and_model_hashes_are_immutable(history):
    _, db, old, _ = history
    with pytest.raises(ValueError, match="cannot be changed"):
        db.update_session(old, {"model2_sha256": "wrong", "grading_config": {"version": "wrong"}})
    snapshot = db.session(old)
    assert snapshot["model2_sha256"] == "old-hash" and snapshot["grading_config"]["version"] == "old"


def test_direct_session_lookup_is_not_limited_to_most_recent_thousand(history):
    client, db, old, _ = history
    with db._lock, db._connection() as connection:
        connection.executemany(
            "INSERT INTO inspection_sessions(session_id, started_at, created_at, updated_at) VALUES (?, ?, ?, ?)",
            [(f"later-{number}", "2025-01-01", "2025-01-01", "2025-01-01") for number in range(1001)],
        )
    assert db.session(old)["model2_sha256"] == "old-hash"
    assert db.session("missing") is None
    response = client.post("/api/export", json={"format": "xlsx", "session_id": old, "fields": ["fish_id"]})
    assert sheet_values(load_workbook(BytesIO(response.content))["Grade Calculation Rules"])["Body original weight"] == 0.7
