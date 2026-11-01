"""Export regression against all seven actual uploaded-fish records."""

from copy import deepcopy
import csv
import hashlib
from io import BytesIO, StringIO
import json
from pathlib import Path

from openpyxl import load_workbook
import pytest

from src.api.export import EXPORT_FIELDS, inspection_row, make_csv, make_xlsx


FIXTURE = Path(__file__).parent / "fixtures" / "fish-inspection-upload-all.csv"
RECORDED_TIMESTAMP = "2026-11-01T00:20:20.811Z"


def reproduction():
    with FIXTURE.open(encoding="utf-8-sig", newline="") as source:
        rows = list(csv.DictReader(source))
    return [json.loads(row["Complete fish event"]) for row in rows], json.loads(rows[0]["Source metadata"])


FIELDS = (
    "fish_id", "timestamp", "grade", "body_width_analysis", "body_width_threshold_configuration",
    "body_width_calibration_status", "body_width_score_status", "body_width_class_c_points",
    "body_width_class_c_points_display", "body_width_score_contributed", "geometric_applied_score_shift",
    "original_weighted_grade", "original_weighted_support", "original_weighted_scores", "final_combined_grade",
    "final_combined_support", "final_combined_scores", "width_changed_classification", "relative_body_width_proxy",
    "body_width_measurement_status", "body_width_reliability_reasons", "body_width_relative_proxy_source",
    "body_width_configuration_source", "whole_fish_roi_width_px", "whole_fish_roi_height_px",
    "body_roi_width_px", "body_roi_height_px", "body_width_roi_proxy", "body_width_detector_proxy",
)


def test_actual_csv_fixture_is_exact_original_and_has_seven_records():
    assert hashlib.sha256(FIXTURE.read_bytes()).hexdigest() == "dc7e06acb4e767650fcb7dc36ff78ab288c6039b6501f5bec8e5c40efded31fd"
    events, _ = reproduction()
    assert len(events) == 7
    assert [event["id"] for event in events] == list(range(1, 8))


def test_all_seven_csv_rows_populate_analysis_calibration_timestamp_and_original_final_grades():
    events, metadata = reproduction()
    before = deepcopy(events)
    encoded = make_csv(events, FIELDS, source_metadata=metadata)
    rows = list(csv.DictReader(StringIO(encoded.decode("utf-8-sig"))))
    assert len(rows) == 7
    for event, row in zip(events, rows):
        assert row[EXPORT_FIELDS["fish_id"]] == f"Fish #{event['id']}"
        assert row[EXPORT_FIELDS["timestamp"]] == RECORDED_TIMESTAMP
        width = json.loads(row[EXPORT_FIELDS["body_width_analysis"]])
        config = json.loads(row[EXPORT_FIELDS["body_width_threshold_configuration"]])
        assert width
        assert config["calibration_validated"] is False
        assert config["scoring_enabled"] is False
        assert row[EXPORT_FIELDS["body_width_calibration_status"]] == "Not Calibrated"
        assert row[EXPORT_FIELDS["body_width_score_status"]] == "Not calibrated"
        assert row[EXPORT_FIELDS["body_width_class_c_points"]] == ""
        assert row[EXPORT_FIELDS["body_width_class_c_points_display"]] == "Not calibrated"
        assert row[EXPORT_FIELDS["body_width_score_contributed"]] == "No"
        assert float(row[EXPORT_FIELDS["geometric_applied_score_shift"]]) == 0
        assert row[EXPORT_FIELDS["original_weighted_grade"]] == event["quality"]
        assert row[EXPORT_FIELDS["final_combined_grade"]] == event["quality"]
        assert json.loads(row[EXPORT_FIELDS["original_weighted_scores"]]) == event["analysis"]["weighted_scores"]
        assert json.loads(row[EXPORT_FIELDS["final_combined_scores"]]) == event["analysis"]["weighted_scores"]
        assert row[EXPORT_FIELDS["width_changed_classification"]] == "No"
        assert "reconstruction" in row[EXPORT_FIELDS["body_width_configuration_source"]]
        expected_support = event["analysis"]["final_score"]
        if expected_support is not None:
            assert float(row[EXPORT_FIELDS["original_weighted_support"]]) == pytest.approx(expected_support * 100, abs=0.0001)
            assert float(row[EXPORT_FIELDS["final_combined_support"]]) == pytest.approx(expected_support * 100, abs=0.0001)
        else:
            assert row[EXPORT_FIELDS["original_weighted_support"]] == ""
            assert row[EXPORT_FIELDS["final_combined_support"]] == ""
    assert events == before
    expected = {1: .386, 2: .481, 3: .265, 5: .302}
    for event, row in zip(events, rows):
        value = row[EXPORT_FIELDS["relative_body_width_proxy"]]
        if event["id"] in expected:
            assert float(value) == pytest.approx(expected[event["id"]], abs=.0006)
            assert float(row[EXPORT_FIELDS["body_width_roi_proxy"]]) == float(value)
            assert row[EXPORT_FIELDS["body_width_relative_proxy_source"]]
            assert row[EXPORT_FIELDS["body_width_measurement_status"]] == "Unreliable"
        else:
            assert value == ""
            assert row[EXPORT_FIELDS["body_width_measurement_status"]] == "Missing Detection"


def test_xlsx_reuses_batch_measurements_and_does_not_drop_timestamp_or_configuration():
    events, metadata = reproduction()
    workbook = load_workbook(BytesIO(make_xlsx(events, FIELDS, {}, source_metadata=metadata)), data_only=True)
    rows = list(workbook["Fish Inspections"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == 7
    for event, values in zip(events, rows):
        row = dict(zip(FIELDS, values))
        assert row["timestamp"] == RECORDED_TIMESTAMP
        assert row["original_weighted_grade"] == row["final_combined_grade"] == event["quality"]
        assert json.loads(row["body_width_analysis"])["threshold_configuration"]
        assert row["body_width_class_c_points"] is None
        assert row["body_width_class_c_points_display"] == "Not calibrated"
        assert row["geometric_applied_score_shift"] == 0


def test_standalone_row_uses_existing_analysis_metadata_timestamp_without_current_time():
    events, metadata = reproduction()
    event = deepcopy(events[0])
    event["analysis"]["inspection_source"] = {"analyzed_at": RECORDED_TIMESTAMP}
    row = inspection_row(event)
    assert row["timestamp"] == RECORDED_TIMESTAMP
    event["timestamp"] = "2026-09-01T12:00:00-04:00"
    assert inspection_row(event, source_metadata=metadata)["timestamp"] == event["timestamp"]


def test_existing_complete_width_result_remains_authoritative_and_bare_history_stays_unknown():
    event = {"track_id": 7, "timestamp": RECORDED_TIMESTAMP, "quality": "Class B",
             "analysis": {"body_width": {"relative_width_proxy": .25, "score_status": "Not calibrated",
                 "threshold_configuration": {"calibration_validated": False}, "score_contributed": False},
                 "weighted_scores": {"Class B": .6}, "final_score": .6}}
    assert json.loads(inspection_row(event)["body_width_analysis"]) == event["analysis"]["body_width"]
    historic = inspection_row({"quality": "Class B", "analysis": {"weighted_scores": {"Class B": .6}, "final_score": .6}})
    assert historic["body_width_analysis"] == ""
    assert historic["original_weighted_grade"] == ""
